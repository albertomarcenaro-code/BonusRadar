"""AI-compiled application documents rendered to PDF (reportlab) and DOCX (python-docx)."""

import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib import colors

from lib.dates import today_iso
from services.llm import ask_json

STORAGE = Path(__file__).parent.parent / "storage" / "documents"

DOC_SYSTEM = (
    "Sei un operatore CAF italiano. Prepari domande di accesso a bonus pubblici precompilate con i dati del "
    "richiedente, in italiano formale, pronte per la firma. Non inventare dati mancanti: lascia '________'."
)


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")[:60] or "documento"


async def compose(bonus: dict, profile: dict) -> dict:
    p = {k: v for k, v in profile.items() if k not in ("_id", "updated_at")}
    prompt = f"""Bonus: {json.dumps({k: bonus.get(k) for k in ('title','authority','amount','deadline','summary','requirements','required_documents','source_url')}, ensure_ascii=False)}
Dati richiedente e nucleo: {json.dumps(p, ensure_ascii=False, default=str)}
Data odierna: {today_iso()}

Restituisci JSON con:
"title": titolo della domanda (es. "Domanda di accesso a ..."),
"addressee": ente destinatario,
"intro": paragrafo "Il/La sottoscritto/a ... CHIEDE ..." compilato,
"sections": array di {{"heading", "fields": [{{"label","value"}}]}} (dati anagrafici, residenza, nucleo familiare, situazione economica, modalità di pagamento IBAN, altri dati specifici del bonus),
"declarations": array di dichiarazioni sostitutive ai sensi del DPR 445/2000 pertinenti ai requisiti,
"attachments": array degli allegati da includere,
"submission_notes": 2-3 frasi su come e dove presentare la domanda (portale, SPID/CIE, scadenza)."""
    return await ask_json(DOC_SYSTEM, prompt)


def _place_date(profile: dict) -> str:
    y, m, d = today_iso().split("-")
    return f"{profile.get('city') or '________'}, {d}/{m}/{y}"


def render_pdf(c: dict, profile: dict, path: Path):
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], fontSize=15, leading=19, alignment=0)
    h2 = ParagraphStyle("h2", parent=ss["Heading3"], textColor=colors.HexColor("#0056B3"), spaceBefore=10)
    body = ParagraphStyle("b", parent=ss["BodyText"], fontSize=10, leading=14)
    esc = lambda s: str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    story = [
        Paragraph(f"Spett.le {esc(c.get('addressee', ''))}", body),
        Spacer(1, 8),
        Paragraph(esc(c.get("title", "Domanda")), h1),
        Paragraph(esc(c.get("intro", "")), body),
    ]
    for sec in c.get("sections", []):
        story.append(Paragraph(esc(sec.get("heading", "")), h2))
        rows = [[Paragraph(f"<b>{esc(f.get('label',''))}</b>", body), Paragraph(esc(f.get("value", "")), body)] for f in sec.get("fields", [])]
        if rows:
            t = Table(rows, colWidths=[6 * cm, 11 * cm])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t)
    if c.get("declarations"):
        story.append(Paragraph("Dichiarazioni sostitutive (DPR 445/2000)", h2))
        story += [Paragraph(f"• {esc(d)}", body) for d in c["declarations"]]
    if c.get("attachments"):
        story.append(Paragraph("Allegati", h2))
        story += [Paragraph(f"☐ {esc(a)}".replace("☐", "[ ]"), body) for a in c["attachments"]]
    story += [
        Spacer(1, 24),
        Paragraph(esc(_place_date(profile)), body),
        Spacer(1, 28),
        Paragraph("Firma del richiedente ______________________________", body),
        Paragraph(esc(profile.get("full_name", "")), body),
    ]
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm).build(story)


def render_docx(c: dict, profile: dict, path: Path):
    d = Document()
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_paragraph(f"Spett.le {c.get('addressee', '')}")
    d.add_heading(c.get("title", "Domanda"), level=1)
    d.add_paragraph(c.get("intro", ""))
    for sec in c.get("sections", []):
        d.add_heading(sec.get("heading", ""), level=2)
        fields = sec.get("fields", [])
        if fields:
            t = d.add_table(rows=0, cols=2)
            t.style = "Table Grid"
            for f in fields:
                row = t.add_row().cells
                row[0].text = str(f.get("label", ""))
                row[1].text = str(f.get("value", ""))
    if c.get("declarations"):
        d.add_heading("Dichiarazioni sostitutive (DPR 445/2000)", level=2)
        for x in c["declarations"]:
            d.add_paragraph(str(x), style="List Bullet")
    if c.get("attachments"):
        d.add_heading("Allegati", level=2)
        for a in c["attachments"]:
            d.add_paragraph(f"[ ] {a}")
    d.add_paragraph("")
    d.add_paragraph(_place_date(profile))
    d.add_paragraph("\nFirma del richiedente ______________________________")
    d.add_paragraph(profile.get("full_name", ""))
    d.save(str(path))


async def generate(folder_id: str, bonus: dict, profile: dict) -> tuple[list[dict], dict]:
    content = await compose(bonus, profile)
    folder = STORAGE / folder_id
    folder.mkdir(parents=True, exist_ok=True)
    base = f"domanda-{slug(bonus['title'])}"
    render_pdf(content, profile, folder / f"{base}.pdf")
    render_docx(content, profile, folder / f"{base}.docx")
    files = [{"name": f"{base}.pdf", "kind": "pdf"}, {"name": f"{base}.docx", "kind": "docx"}]
    return files, content
